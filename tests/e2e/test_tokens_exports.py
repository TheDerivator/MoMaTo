"""API tokens and exports in a real browser: a token is created on the Profile page, shown once,
works for an "agent", marks what it writes "via" and stops working when revoked; the project list
and booked time download as CSV and Excel. [TOK-1] [TOK-3] [TOK-4] [EXP-1] [EXP-2] [EXP-3]"""

import csv
import re
from io import BytesIO, StringIO

from openpyxl import load_workbook
from playwright.sync_api import Page, Playwright, expect

from tests.e2e.conftest import remember, sign_in

SENSORS = 4102


def open_profile(page: Page, base: str, username: str = "tomas.rivera") -> None:
    remember(page)
    page.goto(f"{base}profile")
    sign_in(page, username)
    expect(page.get_by_role("region", name="API tokens")).to_be_visible()


def test_a_token_is_created_shown_once_used_and_revoked(
    sample_server: str, page: Page, playwright: Playwright, console_errors: list[str]
) -> None:
    open_profile(page, sample_server)
    tokens = page.get_by_role("region", name="API tokens")
    expect(tokens).to_contain_text("No tokens yet.")

    tokens.get_by_role("button", name="New token").click()
    dialog = page.get_by_role("dialog", name="New API token")
    dialog.get_by_label("Name").fill("Claude Code")
    dialog.get_by_label(re.compile("Read and write")).check()
    dialog.get_by_label("Expires after").select_option(label="30 days")
    dialog.get_by_role("button", name="Create token").click()

    shown = page.get_by_role("dialog", name="Your new token")
    secret = shown.get_by_test_id("token-secret").inner_text()
    assert re.fullmatch(r"mm_[A-Za-z0-9_-]{40,}", secret)
    expect(shown).to_contain_text(f'$env:MOMATO_TOKEN = "{secret}"')
    with page.expect_download() as download:
        shown.get_by_role("link", name="Download the agent guide").click()
    assert download.value.suggested_filename == "momato-agent-guide.md"
    guide = download.value.path().read_text(encoding="utf-8")
    assert "Authorization: Bearer" in guide
    assert secret not in guide
    shown.get_by_role("button", name="Done").click()

    row = tokens.get_by_role("row", name=re.compile("Claude Code"))
    expect(row).to_contain_text("Read and write")
    expect(row).to_contain_text("Never used")
    expect(row).not_to_contain_text(secret)  # only its start is shown

    # The agent: a client of its own, without the browser's cookies; only the token.
    agent = playwright.request.new_context()
    try:
        headers = {"Authorization": f"Bearer {secret}"}
        written = agent.post(
            f"{sample_server}api/projects/{SENSORS}/remarks",
            headers=headers,
            data={"body_md": "Summary from the agent."},
        )
        assert written.status == 201, written.text()

        page.goto(f"{sample_server}projects/{SENSORS}/remarks")
        remark = page.get_by_role("article", name=re.compile("^Remark by Tomás Rivera")).filter(
            has_text="Summary from the agent."
        )
        expect(remark.locator(".via-badge")).to_have_text("via Claude Code")

        page.goto(f"{sample_server}profile")
        expect(row).to_contain_text("Last used")
        row.get_by_role("button", name="Revoke").click()
        page.get_by_role("dialog").get_by_role("button", name="Revoke").click()
        expect(tokens).to_contain_text("No tokens yet.")
        refused = agent.get(f"{sample_server}api/auth/me", headers=headers)
        assert refused.status == 401
    finally:
        agent.dispose()
    assert console_errors == []


def test_a_read_token_is_refused_when_it_writes(
    sample_server: str, page: Page, playwright: Playwright
) -> None:
    open_profile(page, sample_server)
    page.get_by_role("button", name="New token").click()
    dialog = page.get_by_role("dialog", name="New API token")
    dialog.get_by_label("Name").fill("Reporter")
    dialog.get_by_role("button", name="Create token").click()  # read only is the default
    secret = page.get_by_test_id("token-secret").inner_text()
    page.get_by_role("dialog").get_by_role("button", name="Done").click()
    agent = playwright.request.new_context()
    try:
        headers = {"Authorization": f"Bearer {secret}"}
        assert agent.get(f"{sample_server}api/projects/{SENSORS}", headers=headers).ok
        refused = agent.post(
            f"{sample_server}api/projects/{SENSORS}/remarks",
            headers=headers,
            data={"body_md": "Not allowed."},
        )
        assert refused.status == 403
    finally:
        agent.dispose()


def test_the_project_list_downloads_as_csv_and_excel(sample_server: str, page: Page) -> None:
    remember(page)
    page.goto(f"{sample_server}projects?status=progress")
    sign_in(page, "anna.claes")
    expect(page.locator(".table__link").first).to_be_visible()
    count = int(re.match(r"\d+", page.locator(".list-count").inner_text()).group(0))  # pyright: ignore

    page.get_by_role("button", name="Export").click()
    with page.expect_download() as download:
        page.get_by_role("link", name="CSV (.csv)").click()
    assert download.value.suggested_filename == "projects-2026-10-08.csv"
    text = download.value.path().read_text(encoding="utf-8-sig")
    rows = list(csv.DictReader(StringIO(text)))
    assert len(rows) == count
    assert {r["Status"] for r in rows} == {"progress"}  # the filter of the screen is in force

    with page.expect_download() as excel:  # the popover is still open
        page.get_by_role("link", name="Excel (.xlsx)").click()
    sheet = load_workbook(BytesIO(excel.value.path().read_bytes())).active
    assert sheet is not None
    assert sheet.max_row - 1 == count
    assert sheet.freeze_panes == "A2"


def test_booked_time_downloads_for_a_period_and_branch(sample_server: str, page: Page) -> None:
    remember(page)
    page.goto(f"{sample_server}people?mode=spent&period=week")
    sign_in(page, "anna.claes")
    expect(page.locator(".pp-person").first).to_be_visible()

    page.get_by_role("button", name="Export time").click()
    popover = page.get_by_role("dialog", name="Export booked time")
    expect(popover.get_by_label("From")).to_have_value("2026-09-28")  # what the screen shows
    expect(popover.get_by_label("To")).to_have_value("2026-10-02")
    with page.expect_download() as download:
        popover.get_by_role("link", name="CSV (.csv)").click()
    assert download.value.suggested_filename == "time-entries-2026-09-28-2026-10-02.csv"
    rows = list(csv.DictReader(StringIO(download.value.path().read_text(encoding="utf-8-sig"))))
    assert rows
    assert {r["Week"] for r in rows} == {"2026-W40"}
    assert "Anna Claes" in {r["Person"] for r in rows}

    # a wrong period turns the links off and says why
    popover.get_by_label("To").fill("2026-09-01")
    expect(popover.get_by_role("alert")).to_have_text("The period ends before it starts.")
    expect(popover.get_by_role("link", name="CSV (.csv)")).to_have_count(0)
