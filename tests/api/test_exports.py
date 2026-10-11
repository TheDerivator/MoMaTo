"""Exports: the project list and booked time as CSV and XLSX, with the rights of the caller.
[EXP-1] [EXP-2] [EXP-3] [EXP-4] D-033
"""

import csv
import io
from datetime import date, datetime
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import delete, func, select

from momato.db.models import DefaultGrant, TimeEntry, User
from momato.db.session import Database
from momato.services.sample_data import Sample
from tests.api.conftest import ClientFactory
from tests.api.test_projects import node_id, person_id
from tests.api.test_statement_counts import add_projects

CSV = "text/csv; charset=utf-8"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
LAST_WEEK = {"start": "2026-09-28", "end": "2026-10-02"}  # the "past week" of People planning


def read_csv(response: Any) -> tuple[list[str], list[dict[str, str]]]:
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == CSV
    assert response.content.startswith(b"\xef\xbb\xbf")  # a byte order mark, for Excel
    rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
    header, body = rows[0], rows[1:]
    return header, [dict(zip(header, row, strict=True)) for row in body]


def read_xlsx(response: Any) -> Any:
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == XLSX
    return load_workbook(io.BytesIO(response.content)).active


def listed_numbers(client: TestClient, **params: Any) -> list[int]:
    """The numbers the screen lists for these filters, page after page."""
    numbers: list[int] = []
    after: int | None = None
    while True:
        query = dict(params, limit=100)
        if after is not None:
            query["after"] = after
        page = client.get("/api/projects", params=query).json()
        numbers += [row["number"] for row in page["items"]]
        after = page["next_after"]
        if after is None:
            return numbers


# ---------------------------------------------------------------- who may


def test_a_visitor_gets_401_and_bad_input_422(clients: ClientFactory) -> None:
    visitor, anna = clients(), clients("anna.claes")
    assert visitor.get("/api/exports/projects").status_code == 401
    assert visitor.get("/api/exports/time-entries", params=LAST_WEEK).status_code == 401
    assert anna.get("/api/exports/projects", params={"format": "pdf"}).status_code == 422
    assert anna.get("/api/exports/time-entries").status_code == 422  # a period is needed
    reversed_ = anna.get(
        "/api/exports/time-entries", params={"start": "2026-10-02", "end": "2026-09-28"}
    )
    assert reversed_.status_code == 422
    too_long = anna.get(
        "/api/exports/time-entries", params={"start": "2020-01-01", "end": "2026-10-02"}
    )
    assert too_long.status_code == 422
    far = anna.get("/api/exports/time-entries", params={**LAST_WEEK, "node": 999_999})
    assert far.status_code == 404
    bad_filter = anna.get("/api/exports/projects", params={"node": 999_999})
    assert bad_filter.status_code == 404


# ---------------------------------------------------------------- the project list (EXP-1)


def test_the_csv_holds_the_rows_of_the_screen_with_core_and_detail_fields(
    anna: TestClient, sample: Sample
) -> None:
    header, rows = read_csv(anna.get("/api/exports/projects"))
    assert header[:9] == [
        "Number",
        "Project",
        "Organisation",
        "Status",
        "Priority",
        "Coordinator",
        "Start date",
        "End date",
        "Goal",
    ]
    detail_labels = header[9:]
    assert detail_labels  # every detail field of the profile follows
    assert [int(r["Number"]) for r in rows] == listed_numbers(anna)  # same rows, same order
    first = next(r for r in rows if r["Number"] == "4102")
    assert first["Project"] == "Sensor retrofit on conveyor 4"
    assert first["Organisation"] == "Level 1-2 automation"
    assert first["Coordinator"] == "Tomás Rivera"
    assert first["Status"] == "progress"
    assert first["Start date"] == "2025-11-03"  # ISO dates, no locale
    assert first["End date"] == "2027-03-31"
    assert first["Goal"].startswith("Replace the end-of-life sensors")


def test_the_filters_of_the_screen_narrow_the_export(anna: TestClient) -> None:
    for params in (
        {"status": "progress"},
        {"q": "sensor"},
        {"node": node_id(anna, "IT & automation")},
        {"coordinator": "Rivera", "status": "progress"},
        {"name": "mill"},
    ):
        _, rows = read_csv(anna.get("/api/exports/projects", params=params))
        assert [int(r["Number"]) for r in rows] == listed_numbers(anna, **params), params
    # stopped projects are left out unless their chip is selected (PL-2), as on the screen
    _, default = read_csv(anna.get("/api/exports/projects"))
    assert "stopped" not in {r["Status"] for r in default}
    _, stopped = read_csv(anna.get("/api/exports/projects", params={"status": "stopped"}))
    assert stopped
    assert {r["Status"] for r in stopped} == {"stopped"}


def test_a_filename_says_what_and_when(anna: TestClient) -> None:
    response = anna.get("/api/exports/projects", params={"format": "xlsx"})
    assert (
        response.headers["content-disposition"] == 'attachment; filename="projects-2026-10-08.xlsx"'
    )
    response = anna.get("/api/exports/time-entries", params=LAST_WEEK)
    assert (
        response.headers["content-disposition"]
        == 'attachment; filename="time-entries-2026-09-28-2026-10-02.csv"'
    )


def test_the_xlsx_has_typed_cells_and_a_frozen_header(anna: TestClient) -> None:
    sheet = read_xlsx(anna.get("/api/exports/projects", params={"format": "xlsx"}))
    assert sheet.title == "Projects"
    assert sheet.freeze_panes == "A2"
    header = [c.value for c in sheet[1]]
    assert header[:3] == ["Number", "Project", "Organisation"]
    assert sheet["A1"].font.bold is True
    row = next(r for r in sheet.iter_rows(min_row=2) if r[0].value == 4102)
    assert row[0].data_type == "n"  # a number, not text
    assert isinstance(row[6].value, datetime)  # start date
    assert row[6].value.date() == date(2025, 11, 3)
    assert row[6].number_format == "yyyy-mm-dd"
    assert row[1].value == "Sensor retrofit on conveyor 4"
    # the same rows as the screen
    numbers = [r[0].value for r in sheet.iter_rows(min_row=2)]
    assert numbers == listed_numbers(anna)


def test_a_large_list_is_exported_completely_in_batches(
    clients: ClientFactory, sample_database: Database
) -> None:
    """[D-028] 2,000 projects: more than one batch, nothing repeated or lost."""
    add_projects(sample_database, 1_800)
    anna = clients("anna.claes")
    _, rows = read_csv(anna.get("/api/exports/projects"))
    numbers = [int(r["Number"]) for r in rows]
    assert len(numbers) > 1_900
    assert numbers == sorted(numbers, reverse=True)
    assert len(set(numbers)) == len(numbers)
    assert numbers == listed_numbers(anna)


def test_an_invisible_project_is_not_exported(
    clients: ClientFactory, sample_database: Database, sample: Sample
) -> None:
    """[EXP-4] Nobody is a viewer by default any more: Tomás has only his own projects."""
    with sample_database.session(write=True) as session:
        session.execute(delete(DefaultGrant))
    tomas = clients("tomas.rivera")
    _, rows = read_csv(tomas.get("/api/exports/projects"))
    mine = {p.number for p in sample.projects if "Tomás Rivera" in p.members}
    assert {int(r["Number"]) for r in rows} == {n for n in mine if _shown(sample, n)}


def _shown(sample: Sample, number: int) -> bool:
    project = next(p for p in sample.projects if p.number == number)
    return project.status.value != "stopped"  # the list's default (PL-2)


# ---------------------------------------------------------------- booked time (EXP-2)


def spent_hours(client: TestClient, name: str) -> float:
    body = client.get("/api/people-planning", params={"mode": "spent", "q": name}).json()
    assert body["unit"] == "hours"
    return body["rows"][0]["total"]["value"]


def test_time_entries_equal_the_spent_total_of_the_screen(anna: TestClient) -> None:
    """The past week, as People planning shows it, and the export of the same days agree."""
    header, rows = read_csv(anna.get("/api/exports/time-entries", params=LAST_WEEK))
    assert header == [
        "Date",
        "Week",
        "Person",
        "Organisation",
        "Project number",
        "Project",
        "Task",
        "Category",
        "Hours",
        "Week submitted",
    ]
    mine = [r for r in rows if r["Person"] == "Anna Claes"]
    assert mine
    assert sum(float(r["Hours"]) for r in mine) == pytest.approx(spent_hours(anna, "Anna Claes"))
    assert {r["Week"] for r in mine} == {"2026-W40"}
    assert {r["Organisation"] for r in mine} == {"Quality"}
    assert [r["Date"] for r in rows] == sorted(r["Date"] for r in rows)  # ordered by date
    project_rows = [r for r in mine if r["Project number"]]
    category_rows = [r for r in mine if not r["Project number"]]
    assert project_rows
    assert all(r["Category"] == "" for r in project_rows)
    assert all(r["Project"] == "" and r["Category"] for r in category_rows)
    assert {r["Week submitted"] for r in mine} <= {"yes", "no"}


def test_a_branch_narrows_the_time_export(anna: TestClient) -> None:
    quality = node_id(anna, "Quality")
    _, rows = read_csv(anna.get("/api/exports/time-entries", params={**LAST_WEEK, "node": quality}))
    assert rows
    assert {r["Organisation"] for r in rows} == {"Quality"}
    everyone = read_csv(anna.get("/api/exports/time-entries", params=LAST_WEEK))[1]
    assert len(everyone) > len(rows)


def test_the_xlsx_of_time_has_numbers_dates_and_booleans(anna: TestClient) -> None:
    response = anna.get("/api/exports/time-entries", params={**LAST_WEEK, "format": "xlsx"})
    sheet = read_xlsx(response)
    assert sheet.title == "Time entries"
    assert sheet.freeze_panes == "A2"
    row = next(r for r in sheet.iter_rows(min_row=2))
    assert isinstance(row[0].value, datetime)
    assert row[8].data_type == "n"  # hours
    assert isinstance(row[8].value, int | float)
    assert row[9].data_type == "b"  # week submitted
    csv_rows = read_csv(anna.get("/api/exports/time-entries", params=LAST_WEEK))[1]
    assert sheet.max_row - 1 == len(csv_rows)


def test_someone_without_hours_view_exports_only_their_own_time(
    clients: ClientFactory, sample_database: Database
) -> None:
    """[EXP-4] Without the default Hours viewer on their department: just their own hours."""
    with sample_database.session(write=True) as session:
        session.execute(delete(DefaultGrant).where(DefaultGrant.scope_kind == "own_department"))
    anna = clients("anna.claes")
    _, rows = read_csv(anna.get("/api/exports/time-entries", params=LAST_WEEK))
    assert rows
    assert {r["Person"] for r in rows} == {"Anna Claes"}
    # ... and a colleague's export holds that colleague only
    bram = clients("bram.peeters")
    assert {
        r["Person"] for r in read_csv(bram.get("/api/exports/time-entries", params=LAST_WEEK))[1]
    } == {"Bram Peeters"}


def test_the_export_holds_exactly_what_the_database_holds_for_the_period(
    admin: TestClient, sample_database: Database
) -> None:
    with sample_database.session() as session:
        count = session.scalar(
            select(func.count())
            .select_from(TimeEntry)
            .join(User, User.id == TimeEntry.user_id)
            .where(TimeEntry.work_date.between(date(2026, 9, 28), date(2026, 10, 2)))
        )
    _, rows = read_csv(admin.get("/api/exports/time-entries", params=LAST_WEEK))
    assert len(rows) == count
    assert person_id(admin, "Anna Claes")  # the export names people, never ids


# ---------------------------------------------------------------- with a token


def test_a_read_token_exports_what_its_owner_sees(clients: ClientFactory, app: Any) -> None:
    """[TOK-1] Exports are reads: a read-only token may use them."""
    owner = clients("anna.claes")
    secret = owner.post("/api/auth/tokens", json={"name": "Reporter"}).json()["secret"]
    agent = TestClient(app)
    response = agent.get(
        "/api/exports/time-entries", params=LAST_WEEK, headers={"Authorization": f"Bearer {secret}"}
    )
    assert (
        read_csv(response)[1]
        == read_csv(owner.get("/api/exports/time-entries", params=LAST_WEEK))[1]
    )
